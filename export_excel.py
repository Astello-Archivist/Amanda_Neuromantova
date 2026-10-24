"""Create a styled, sortable workbook from the catalog CSV."""
from pathlib import Path
from collections import Counter
from datetime import datetime
from zoneinfo import ZoneInfo
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parent
NAVY, GREEN, PALE, AMBER = '172B35', '2D6A4F', 'EAF4EF', 'FFF2CC'


def literal(cell, value):
    cell.value = str(value or '')[:32767]
    cell.data_type = 's'


def main():
    with (ROOT/'tracks.csv').open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f, delimiter=';'))
    wb = Workbook()
    summary = wb.active
    summary.title = 'Обзор'
    summary.sheet_view.showGridLines = False
    summary.column_dimensions['A'].width = 4
    summary.column_dimensions['B'].width = 38
    summary.column_dimensions['C'].width = 65
    summary.merge_cells('B2:F3')
    summary['B2'] = 'CH—D  /  МУЗЫКАЛЬНЫЙ АРХИВ'
    summary['B2'].font = Font(size=24, bold=True, color='FFFFFF')
    summary['B2'].fill = PatternFill('solid', fgColor=NAVY)
    metrics = [('Всего треков', len(rows)), ('Со ссылками на прослушивание', sum(any(r.get(k) for k in ('audius_url', 'youtube_url', 'drive_url')) for r in rows)),
               ('Со встроенной обложкой', sum(bool(r['cover_path']) for r in rows)),
               ('Без встроенной обложки', sum(not r['cover_path'] for r in rows)),
               ('С текстом песни', sum(bool(r['lyrics']) for r in rows))]
    for n, (label, value) in enumerate(metrics, 5):
        summary.cell(n, 2, label).font = Font(size=13, color=NAVY)
        summary.cell(n, 3, value).font = Font(size=20, bold=True, color=GREEN)
    instructions = [
        'Лист «Треки»: фильтры в заголовках, сортировка, закреплённые название и исполнитель.',
        'Жёлтые ячейки YouTube — ссылка ещё не сопоставлена. Пустые теги ещё не заполнены.',
        'Лист «Без обложек»: песни, для которых в MP4 использована заглушка.',
        'Лист «Тексты»: только реально найденные тексты; жанр из метаданных не означает анализ звучания.',
        'Лист «Файлы»: технические поля и SHA-256 для поиска дубликатов.',
        'Этот Excel — снимок tracks.csv. Правки Excel автоматически не переносятся в CSV и веб-каталог.',
        'Для повторного создания: python music-catalog/export_excel.py',
    ]
    for n, text in enumerate(instructions, 12):
        summary.merge_cells(start_row=n, start_column=2, end_row=n, end_column=6)
        summary.cell(n, 2, text).alignment = Alignment(wrap_text=True, vertical='center')
        summary.row_dimensions[n].height = 36
    columns = [('ID', 'id', 15), ('Название', 'title_draft', 42), ('Исполнитель', 'artist_draft', 28),
               ('Персонажи', 'characters', 28), ('Предметы', 'items', 26), ('Именные животные', 'named_animals', 25),
               ('Жанр', 'genre', 25), ('Язык', 'language', 16), ('YouTube', 'youtube_url', 40),
               ('Google Drive', 'drive_url', 35), ('Audius', 'audius_url', 45), ('Длительность', '_time', 17),
               ('Обложка', '_cover', 25), ('Альбом', 'metadata_album', 28), ('Примечания', 'notes', 45),
               ('Добавлен в каталог', 'added_at', 23), ('Исходный файл', 'file', 48)]

    def sheet(name, selected, fields, table_name):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = GREEN
        end = get_column_letter(len(fields))
        ws.merge_cells(f'A1:{end}2')
        ws['A1'] = f'CH—D  /  {name.upper()}  ·  {len(selected)}'
        ws['A1'].font = Font(size=20, bold=True, color='FFFFFF')
        ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
        ws['A3'] = '← Обзор'
        ws['A3'].hyperlink = "#'Обзор'!B2"
        ws['A3'].font = Font(color=GREEN, underline='single')
        for col, (label, key, width) in enumerate(fields, 1):
            cell = ws.cell(5, col, label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=GREEN)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[5].height = 32
        for index, r in enumerate(selected, 6):
            for col, (_, key, _) in enumerate(fields, 1):
                c = ws.cell(index, col)
                value = r.get(key, '')
                if key == 'genre': value = r.get('genre') or r.get('metadata_genre', '')
                if key == 'artist_draft': value = r.get('artist_draft') or r.get('metadata_artist', '')
                if key == '_time':
                    c.value = float(r.get('duration_seconds') or 0)/86400
                    c.number_format = '[m]:ss'
                elif key == 'added_at' and value:
                    c.value = datetime.fromisoformat(value).astimezone(ZoneInfo('Europe/Moscow')).replace(tzinfo=None)
                    c.number_format = 'dd.mm.yyyy hh:mm'
                else:
                    if key == '_cover': value = 'Встроенная' if r['cover_path'] else 'Заглушка в MP4'
                    literal(c, value)
                c.font = Font(size=11, color=NAVY)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                if key in ('youtube_url', 'drive_url', 'audius_url'):
                    if value.startswith(('https://', 'http://')):
                        c.hyperlink = value
                        c.font = Font(color='146CB4', underline='single')
                    elif key == 'youtube_url': c.fill = PatternFill('solid', fgColor=AMBER)
                if key == '_cover':
                    c.fill = PatternFill('solid', fgColor=PALE if r['cover_path'] else AMBER)
            ws.row_dimensions[index].height = 43
        if selected:
            table = Table(displayName=table_name, ref=f'A5:{end}{5+len(selected)}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True)
            ws.add_table(table)
        ws.freeze_panes = 'D6' if len(fields)>3 else 'A6'
        ws.sheet_view.zoomScale = 85
        ws.print_title_rows = '1:5'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return ws

    sheet('Треки', rows, columns, 'Tracks')
    sheet('Без обложек', [r for r in rows if not r['cover_path']], columns, 'MissingCovers')
    sheet('Тексты', [r for r in rows if r['lyrics']], [('ID','id',15),('Название','title_draft',40),('Текст','lyrics',100)], 'Lyrics')
    technical = [('ID','id',15),('Файл','file',48),('Название в метаданных','metadata_title',45),
                 ('Исполнитель в метаданных','metadata_artist',30),('Жанр в метаданных','metadata_genre',25),
                 ('Год','metadata_year',14),('Битрейт, кбит/с','bitrate_kbps',20),('Размер, байт','size_bytes',22),
                 ('SHA-256','sha256',70),('Путь обложки','cover_path',35)]
    sheet('Файлы', rows, technical, 'Files')
    summary['B21'] = 'Открыть каталог треков →'
    summary['B21'].hyperlink = "#'Треки'!A1"
    summary['B21'].font = Font(size=14, bold=True, color=GREEN)
    path = ROOT/'Ch-D_music_catalog.xlsx'
    try:
        wb.save(path)
    except PermissionError:
        path = ROOT/'Ch-D_music_catalog_updated.xlsx'
        wb.save(path)
        print('Original workbook is open; saved Ch-D_music_catalog_updated.xlsx')
    check = load_workbook(path)
    assert check['Треки'].max_row == len(rows)+5
    assert len(check['Треки'].tables) == 1
    print(f'Excel created: {len(rows)} tracks, {len(wb.sheetnames)} sheets')


if __name__ == '__main__':
    main()
